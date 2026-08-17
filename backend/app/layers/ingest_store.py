"""
Layer 2 — Ingest & Store (capture and historize).

Production stack:
  Redpanda / Kafka   durable ordered event log
  TimescaleDB        time-series historian (Bronze, operational)
  MinIO              S3-compatible data lake (Medallion Bronze → Silver → Gold)
  DuckDB + Parquet   lake transforms and queries

This laptop reference implements the same contracts:
  hop 1  consume Connectivity's MQTT-confirmed reading → produce scada.telemetry
  hop 2  TimescaleDB INSERT of the raw Kafka payload (Bronze historian)
  hop 3  MinIO landing of the same payload as immutable Parquet (Bronze lake)
  hop 4  DuckDB reads Bronze Parquet, conforms types/units/quality → Silver
  hop 5  DuckDB curates Silver → Gold lake tables for Contextualize (Flink)

Medallion Gold here is a curated telemetry table. It is NOT Layer 3's
Flink-contextualized Gold events (those join MES / SAP / asset model).
"""
from __future__ import annotations

import json
import os
from collections import deque
from datetime import datetime, timezone
from pathlib import Path
from threading import Lock
from typing import Any, Deque, Dict, List, Optional

from ..domain import PARAMETERS, TAG_TO_PARAM
from . import floor_db

LAKE_ROOT = Path(__file__).resolve().parents[2] / "data" / "lake"
KAFKA_TOPIC = "scada.telemetry"
KAFKA_PARTITION = 0


def _now_iso() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def _duck_path(p: Path) -> str:
    return p.resolve().as_posix()


def _jsonable(v: Any) -> Any:
    if v is None:
        return None
    if hasattr(v, "isoformat"):
        return v.isoformat()
    if isinstance(v, bytes):
        return v.decode("utf-8", errors="replace")
    if isinstance(v, (str, int, float, bool)):
        return v
    return str(v)


def _row_jsonable(row: Dict) -> Dict:
    return {k: _jsonable(v) for k, v in row.items()}


def _param_of(reading: Dict) -> Optional[str]:
    if reading.get("param") in PARAMETERS:
        return reading["param"]
    return TAG_TO_PARAM.get(reading.get("tag") or "")


def _canonical_unit(reading: Dict) -> str:
    pid = _param_of(reading)
    if pid:
        return PARAMETERS[pid].unit
    return reading.get("unit") or ""


def _canonical_asset(reading: Dict) -> str:
    if reading.get("asset"):
        return str(reading["asset"])
    pid = _param_of(reading)
    if pid:
        return PARAMETERS[pid].asset
    return ""


def _sql_str(v: Any) -> str:
    return str(v or "").replace("'", "''")


def _catalog_case(tag_col: str, attr: str) -> str:
    parts = []
    for p in PARAMETERS.values():
        parts.append(f"WHEN '{p.tag}' THEN '{_sql_str(getattr(p, attr))}'")
    return f"CASE {tag_col} {' '.join(parts)} ELSE '' END"


def _kafka_value(reading: Dict) -> Dict:
    return {
        "tag": reading.get("tag"),
        "value": reading.get("value"),
        "unit": reading.get("unit"),
        "timestamp": reading.get("timestamp") or reading.get("file_ts"),
        "quality": reading.get("quality") or "GOOD",
        "param": _param_of(reading),
        "asset": reading.get("asset") or _canonical_asset(reading),
        "name": reading.get("name"),
        "uns_topic": reading.get("uns_topic"),
        "source": reading.get("source"),
        "filename": reading.get("filename"),
        "file_ts": reading.get("file_ts"),
        "reading_id": reading.get("reading_id"),
        "row_index": reading.get("row_index"),
    }


class Store:
    def __init__(self, maxlen: int = 5000):
        self._lock = Lock()
        self.historian: Deque[Dict] = deque(maxlen=maxlen)
        self.gold: Deque[Dict] = deque(maxlen=maxlen)
        self.gold_lake: Deque[Dict] = deque(maxlen=maxlen)
        self.kafka_log: Deque[Dict] = deque(maxlen=maxlen)
        self._offset = 8400
        self.steps_done = 0
        self.last_reading: Optional[Dict] = None
        self.last_readings: List[Dict] = []
        self.last_kafka: Optional[Dict] = None
        self.last_kafka_batch: List[Dict] = []
        self.bronze_path: Optional[str] = None
        self.silver_path: Optional[str] = None
        self.gold_path: Optional[str] = None
        self.historian_backend: str = "memory"

    def next_offset(self) -> int:
        with self._lock:
            self._offset += 1
            return self._offset

    def write_bronze(self, reading: Dict, offset: Optional[int] = None) -> int:
        if offset is None:
            offset = self.next_offset()
        rec = {**reading, "kafka_topic": KAFKA_TOPIC, "offset": offset,
               "zone": "bronze"}
        with self._lock:
            self.historian.append(rec)
        return offset

    def write_gold(self, event: Dict) -> None:
        with self._lock:
            self.gold.append({**event, "zone": "gold"})

    def recent_historian(self, n: int = 50) -> List[Dict]:
        with self._lock:
            return list(self.historian)[-n:][::-1]

    def recent_gold(self, n: int = 50, param: str | None = None) -> List[Dict]:
        with self._lock:
            items = list(self.gold)
        if param:
            items = [e for e in items if e.get("param") == param]
        return items[-n:][::-1]

    def stats(self) -> Dict:
        with self._lock:
            return {
                "historian_count": len(self.historian),
                "gold_count": len(self.gold),
                "gold_lake_count": len(self.gold_lake),
                "latest_offset": self._offset,
                "ingest_steps_done": self.steps_done,
                "bronze_path": self.bronze_path,
                "silver_path": self.silver_path,
                "gold_path": self.gold_path,
                "historian_backend": self.historian_backend,
            }

    def reset_steps(self) -> Dict:
        with self._lock:
            self.steps_done = 0
            self.last_kafka = None
            self.last_kafka_batch = []
            self.bronze_path = None
            self.silver_path = None
            self.gold_path = None
        return self.status()

    def purge(self) -> Dict:
        """Wipe in-memory historian/lake and files under data/lake."""
        with self._lock:
            self.historian.clear()
            self.gold.clear()
            self.gold_lake.clear()
            self.kafka_log.clear()
            self.last_reading = None
            self.last_readings = []
            self.last_kafka = None
            self.last_kafka_batch = []
            self.bronze_path = None
            self.silver_path = None
            self.gold_path = None
            self.steps_done = 0
        removed = 0
        if LAKE_ROOT.exists():
            for p in LAKE_ROOT.rglob("*"):
                if p.is_file():
                    try:
                        p.unlink()
                        removed += 1
                    except OSError:
                        pass
        st = self.status()
        st["purged_files"] = removed
        return st

    def status(self) -> Dict:
        with self._lock:
            kafka = self.last_kafka
            return {
                "steps_done": self.steps_done,
                "topic": KAFKA_TOPIC,
                "latest_offset": self._offset,
                "has_reading": bool(self.last_readings or self.last_reading or kafka),
                "reading": self.last_reading,
                "count": len(self.last_readings or self.last_kafka_batch or []),
                "kafka": {
                    "topic": kafka.get("topic") if kafka else None,
                    "offset": kafka.get("offset") if kafka else None,
                    "key": kafka.get("key") if kafka else None,
                } if kafka else None,
                "bronze_path": self.bronze_path,
                "silver_path": self.silver_path,
                "gold_path": self.gold_path,
                "historian_count": len(self.historian),
                "gold_lake_count": len(self.gold_lake),
                "historian_backend": self.historian_backend,
                "lake_root": str(LAKE_ROOT),
            }

    def run_pipeline(self, reading: Optional[Dict] = None,
                     readings: Optional[List[Dict]] = None) -> Dict:
        """Journey / API shortcut: run hops 1–5 on the Connectivity batch."""
        self.reset_steps()
        last = None
        for n in range(1, 6):
            last = self.run_step(n, reading=reading, readings=readings)
        offset = last.get("offset") if last else None
        return {
            "ingested": True,
            "offset": offset,
            "topic": KAFKA_TOPIC,
            "count": last.get("count") if last else 0,
            "bronze_path": self.bronze_path,
            "silver_path": self.silver_path,
            "gold_path": self.gold_path,
            "steps_done": self.steps_done,
        }

    def _coerce_batch(self, reading: Optional[Dict],
                      readings: Optional[List[Dict]]) -> List[Dict]:
        batch: List[Dict] = []
        if readings:
            batch = [r for r in readings if r and r.get("tag") is not None]
        elif reading and reading.get("tag") is not None:
            batch = [reading]
        elif self.last_readings:
            batch = list(self.last_readings)
        elif self.last_reading:
            batch = [self.last_reading]
        if not batch:
            raise ValueError(
                "No Connectivity readings to ingest. Finish Connectivity hop 4 "
                "(subscribe and confirm) so the MQTT stream is on the wire, "
                "then return here."
            )
        return batch

    def run_step(self, step: int, reading: Optional[Dict] = None,
                 readings: Optional[List[Dict]] = None) -> Dict:
        if step < 1 or step > 5:
            raise ValueError("Ingest hops are 1–5")
        batch = self._coerce_batch(reading, readings)
        if step > 1 and self.steps_done < step - 1:
            raise ValueError(f"Complete hop {step - 1} first — hops are serial.")
        if step == 1:
            return self._step_kafka(batch)
        if step == 2:
            return self._step_timescale()
        if step == 3:
            return self._step_minio_bronze()
        if step == 4:
            return self._step_duckdb_silver()
        return self._step_gold_lake()

    # ---- hop 1: Redpanda / Kafka ------------------------------------------
    def _step_kafka(self, batch: List[Dict]) -> Dict:
        envelopes: List[Dict] = []
        log_path = LAKE_ROOT / "kafka" / f"{KAFKA_TOPIC}.jsonl"
        log_path.parent.mkdir(parents=True, exist_ok=True)
        with log_path.open("a", encoding="utf-8") as fh:
            for reading in batch:
                offset = self.next_offset()
                envelope = {
                    "topic": KAFKA_TOPIC,
                    "partition": KAFKA_PARTITION,
                    "offset": offset,
                    "key": reading.get("tag"),
                    "timestamp": _now_iso(),
                    "headers": {
                        "content-type": "application/json",
                        "source": "mqtt-uns",
                        "from_layer": "connectivity",
                    },
                    "value": _kafka_value(reading),
                }
                fh.write(json.dumps(envelope, default=str) + "\n")
                envelopes.append(envelope)
                with self._lock:
                    self.kafka_log.append(envelope)
        first_off = envelopes[0]["offset"]
        last_off = envelopes[-1]["offset"]
        tags = sorted({e["key"] for e in envelopes})
        with self._lock:
            self.last_kafka_batch = envelopes
            self.last_kafka = envelopes[-1]
            self.last_readings = list(batch)
            self.last_reading = dict(batch[-1])
            self.steps_done = 1
            self.bronze_path = None
            self.silver_path = None
            self.gold_path = None
        produce_lines = [
            f"[kafka]    produce  key={e['key']}  offset={e['offset']}  "
            f"{e['value'].get('tag')}={e['value'].get('value')} {e['value'].get('unit') or ''}"
            for e in envelopes[:12]
        ]
        if len(envelopes) > 12:
            produce_lines.append(f"[kafka]    … {len(envelopes) - 12} more messages")
        log = [
            "[redpanda] kafka-compatible broker  127.0.0.1:9092  (topic auto-create)",
            f"[kafka]    produce  topic={KAFKA_TOPIC}  partition={KAFKA_PARTITION}  "
            f"{len(envelopes)} messages  offsets {first_off}–{last_off}",
            f"[kafka]    keys     {', '.join(str(t) for t in tags)}",
            "[kafka]    header  source=mqtt-uns  from_layer=connectivity",
        ] + produce_lines + [
            f"[wal]      append {log_path.as_posix()}",
            "[role]     Connectivity proved the floor is on MQTT. Kafka is the ingest backbone:",
            "[role]     ordered, replayable, fan-out to historian AND the lake.",
            f"✓ produced {len(envelopes)} readings → {KAFKA_TOPIC}@{first_off}–{last_off}",
        ]
        return {
            "step": 1, "ok": True, "offset": last_off, "topic": KAFKA_TOPIC,
            "partition": KAFKA_PARTITION, "count": len(envelopes),
            "offsets": [e["offset"] for e in envelopes],
            "reading": batch[-1], "readings": batch, "kafka": envelopes[-1],
            "log": log,
            "headline": f"Produced {len(envelopes)} messages → {KAFKA_TOPIC} offsets {first_off}–{last_off}",
            "steps_done": 1,
        }

    # ---- hop 2: TimescaleDB Bronze historian ------------------------------
    def _step_timescale(self) -> Dict:
        batch = self.last_kafka_batch
        if not batch:
            raise ValueError("Produce to Kafka first (hop 1).")
        recs = []
        for msg in batch:
            val = msg["value"]
            offset = msg["offset"]
            ts_raw = val.get("timestamp") or msg.get("timestamp")
            recs.append({
                **val,
                "kafka_offset": offset,
                "ts_raw": ts_raw,
            })
            self.write_bronze({k: v for k, v in val.items() if k}, offset=offset)
        backend, pg_log = self._historian_insert(recs)
        self.historian_backend = backend
        with self._lock:
            self.steps_done = max(self.steps_done, 2)
        first_off = batch[0]["offset"]
        last_off = batch[-1]["offset"]
        tags = sorted({r.get("tag") for r in recs})
        log = [
            "[timescaledb] Bronze historian — operational time-series store",
            "[timescaledb] CREATE TABLE IF NOT EXISTS historian_bronze (...);",
            "[timescaledb] SELECT create_hypertable('historian_bronze', 'ts', if_not_exists => TRUE);",
            f"[consumer]   read {KAFKA_TOPIC} offsets {first_off}–{last_off}  ({len(recs)} rows)",
            f"[sql]        INSERT INTO historian_bronze (…) VALUES (…) × {len(recs)}",
            f"[sql]        tags {', '.join(str(t) for t in tags)}",
        ] + pg_log + [
            "[role]        Raw as-is. No unit cleanup, no quality drop, no MES join.",
            "[role]        This Bronze is for last-N / by-tag queries on the plant floor.",
            f"✓ historian INSERT {len(recs)} rows  offsets {first_off}–{last_off}  ({backend})",
        ]
        return {
            "step": 2, "ok": True, "offset": last_off, "topic": KAFKA_TOPIC,
            "zone": "bronze", "store": "timescaledb", "count": len(recs),
            "backend": backend, "reading": self.last_reading, "log": log,
            "headline": f"Bronze historian ← {len(recs)} rows offsets {first_off}–{last_off}",
            "steps_done": self.steps_done,
        }

    def _historian_insert(self, recs: List[Dict]) -> tuple[str, List[str]]:
        ddl = """
        CREATE TABLE IF NOT EXISTS historian_bronze (
            kafka_offset BIGINT NOT NULL,
            ts TIMESTAMPTZ,
            ts_raw TEXT,
            tag TEXT NOT NULL,
            value DOUBLE PRECISION NOT NULL,
            unit TEXT,
            quality TEXT,
            uns_topic TEXT,
            source TEXT,
            param TEXT,
            asset TEXT,
            ingested_at TIMESTAMPTZ NOT NULL DEFAULT NOW(),
            PRIMARY KEY (kafka_offset)
        );
        """
        insert_sql = """
        INSERT INTO historian_bronze
            (kafka_offset, ts, ts_raw, tag, value, unit, quality, uns_topic, source, param, asset)
        VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
        ON CONFLICT (kafka_offset) DO NOTHING;
        """
        logs: List[str] = []
        try:
            import psycopg2
            dsn = floor_db.dsn()
            with psycopg2.connect(dsn, connect_timeout=3) as conn:
                conn.autocommit = True
                with conn.cursor() as cur:
                    cur.execute(ddl)
                    hyper = False
                    try:
                        cur.execute("CREATE EXTENSION IF NOT EXISTS timescaledb")
                        cur.execute(
                            "SELECT create_hypertable('historian_bronze', 'ts', "
                            "if_not_exists => TRUE)"
                        )
                        hyper = True
                    except Exception as ext_err:
                        logs.append(
                            f"[timescaledb] extension not loaded on this Postgres "
                            f"({ext_err.__class__.__name__}) — same INSERT contract, regular table"
                        )
                    for rec in recs:
                        ts_raw = rec.get("ts_raw")
                        cur.execute(insert_sql, (
                            rec.get("kafka_offset"),
                            ts_raw, ts_raw,
                            rec.get("tag"), rec.get("value"), rec.get("unit"),
                            rec.get("quality"), rec.get("uns_topic"), rec.get("source"),
                            rec.get("param"), rec.get("asset"),
                        ))
                    cur.execute("SELECT COUNT(*) FROM historian_bronze")
                    n_all = cur.fetchone()[0]
            host = os.environ.get("MOSAIC_PG_HOST", "127.0.0.1")
            port = os.environ.get("MOSAIC_PG_PORT", "5434")
            kind = "TimescaleDB hypertable" if hyper else "Postgres historian table"
            logs.append(f"[timescaledb] connected {host}:{port}/{os.environ.get('MOSAIC_PG_DB', 'mosaic')}")
            logs.append(f"[timescaledb] {kind}  inserted {len(recs)}  table rows now {n_all}")
            return ("timescaledb" if hyper else "postgres", logs)
        except Exception as e:
            logs.append(
                f"[timescaledb] live DB unreachable ({e.__class__.__name__}: {e}) "
                "— in-memory historian still holds the Bronze rows"
            )
            logs.append("[timescaledb] production: docker compose --profile l2 up timescaledb")
            return ("memory", logs)

    # ---- hop 3: MinIO Bronze lake -----------------------------------------
    def _step_minio_bronze(self) -> Dict:
        batch = self.last_kafka_batch
        if not batch:
            raise ValueError("Produce to Kafka first (hop 1).")
        first_off = batch[0]["offset"]
        last_off = batch[-1]["offset"]
        day = _now_iso()[:10]
        rows = []
        for msg in batch:
            val = dict(msg["value"])
            offset = msg["offset"]
            ts_raw = val.get("timestamp") or msg.get("timestamp") or _now_iso()
            if ts_raw:
                day = str(ts_raw)[:10] or day
            rows.append({
                "kafka_offset": int(offset),
                "kafka_topic": KAFKA_TOPIC,
                "kafka_partition": int(KAFKA_PARTITION),
                "tag": str(val.get("tag") or "unknown"),
                "value": float(val["value"]) if val.get("value") is not None else None,
                "unit": val.get("unit") or "",
                "ts_raw": str(ts_raw),
                "quality": val.get("quality") or "GOOD",
                "param": val.get("param") or "",
                "asset": val.get("asset") or "",
                "name": val.get("name") or "",
                "uns_topic": val.get("uns_topic") or "",
                "source": val.get("source") or "",
                "filename": val.get("filename") or "",
                "medallion_zone": "bronze",
            })
        rel = Path("bronze") / f"dt={day}" / f"batch-{first_off}-{last_off}.parquet"
        dest = LAKE_ROOT / rel
        s3_uri = f"s3://mosaic-lake/{rel.as_posix()}"
        import duckdb
        con = duckdb.connect()
        self._register_rows(con, "bronze_row", rows)
        dest.parent.mkdir(parents=True, exist_ok=True)
        con.execute(f"COPY bronze_row TO '{_duck_path(dest)}' (FORMAT PARQUET)")
        preview = self._fetch_dicts(con, "SELECT * FROM bronze_row")
        n = con.execute("SELECT COUNT(*) FROM bronze_row").fetchone()[0]
        con.close()
        with self._lock:
            self.bronze_path = str(dest)
            self.steps_done = max(self.steps_done, 3)
        log = [
            "[minio]   S3-compatible lake  http://127.0.0.1:9000  bucket mosaic-lake",
            "[minio]   dual-write: hop 2 was the historian; hop 3 is the immutable landing zone",
            f"[minio]   PUT  {s3_uri}  ({n} rows)",
            f"[duckdb]  COPY bronze_row TO '{dest.as_posix()}' (FORMAT PARQUET)",
            "[schema]  as-is Kafka payload + kafka_offset — no casts, no filters",
            "[role]    Bronze lake is replayable. Rebuild Silver/Gold from these objects.",
            f"✓ bronze parquet  {n} rows  {s3_uri}",
        ]
        return {
            "step": 3, "ok": True, "offset": last_off, "zone": "bronze",
            "store": "minio", "path": str(dest), "s3_uri": s3_uri, "count": n,
            "preview": preview[:25], "reading": self.last_reading, "log": log,
            "headline": f"Bronze lake ← {n} rows {s3_uri}",
            "steps_done": self.steps_done,
        }

    # ---- hop 4: DuckDB Silver ---------------------------------------------
    def _step_duckdb_silver(self) -> Dict:
        if not self.bronze_path:
            raise ValueError("Write Bronze parquet first (hop 3).")
        batch = self.last_kafka_batch or []
        first_off = batch[0]["offset"] if batch else 0
        last_off = batch[-1]["offset"] if batch else 0
        day = _now_iso()[:10]
        rel = Path("silver") / f"dt={day}" / f"batch-{first_off}-{last_off}.parquet"
        dest = LAKE_ROOT / rel
        s3_uri = f"s3://mosaic-lake/{rel.as_posix()}"
        bronze_uri = _duck_path(Path(self.bronze_path))
        dest.parent.mkdir(parents=True, exist_ok=True)
        param_case = _catalog_case("tag", "id")
        asset_case = _catalog_case("tag", "asset")
        unit_case = _catalog_case("tag", "unit")
        silver_sql = f"""
CREATE OR REPLACE TABLE silver AS
SELECT
    tag,
    COALESCE(NULLIF(param, ''), {param_case})         AS param,
    COALESCE(NULLIF(asset, ''), {asset_case})         AS asset,
    CAST(value AS DOUBLE)                             AS value,
    {unit_case}                                       AS unit,
    TRY_CAST(ts_raw AS TIMESTAMP)                     AS ts,
    UPPER(COALESCE(NULLIF(quality, ''), 'GOOD'))      AS quality,
    uns_topic,
    source,
    kafka_offset                                      AS bronze_offset,
    'silver'                                          AS medallion_zone,
    1                                                 AS schema_version
FROM read_parquet('{bronze_uri}')
WHERE value IS NOT NULL
  AND UPPER(COALESCE(quality, 'GOOD')) <> 'BAD'
"""
        import duckdb
        con = duckdb.connect()
        con.execute(silver_sql)
        con.execute(f"COPY silver TO '{_duck_path(dest)}' (FORMAT PARQUET)")
        preview = self._fetch_dicts(con, "SELECT * FROM silver")
        n = con.execute("SELECT COUNT(*) FROM silver").fetchone()[0]
        n_tags = con.execute("SELECT COUNT(DISTINCT tag) FROM silver").fetchone()[0]
        bronze_cols = [r["column_name"] for r in self._fetch_dicts(
            con, f"DESCRIBE SELECT * FROM read_parquet('{bronze_uri}')")]
        silver_cols = [r["column_name"] for r in self._fetch_dicts(con, "DESCRIBE silver")]
        con.close()
        with self._lock:
            self.silver_path = str(dest)
            self.steps_done = max(self.steps_done, 4)
        log = [
            "[duckdb]  in-process SQL engine over Parquet — no Spark cluster required",
            f"[duckdb]  READ  bronze  {bronze_uri}",
            "[transform] CAST value → DOUBLE  (every row)",
            "[transform] unit → canonical from parameter catalog (per tag)",
            "[transform] TRY_CAST ts_raw → TIMESTAMP",
            "[transform] quality UPPER(); drop BAD rows",
            "[transform] add bronze_offset, medallion_zone, schema_version",
            f"[duckdb]  {silver_sql.strip()}",
            f"[minio]   PUT  {s3_uri}  ({n} rows, {n_tags} tags)",
            f"[schema]  bronze columns: {', '.join(bronze_cols)}",
            f"[schema]  silver columns: {', '.join(silver_cols)}",
            "[role]    Silver is conformed and analytics-ready. Still no MES/SAP/spec.",
            f"✓ silver parquet  {n} rows  {s3_uri}",
        ]
        return {
            "step": 4, "ok": True, "offset": last_off, "zone": "silver",
            "store": "duckdb", "path": str(dest), "s3_uri": s3_uri, "count": n,
            "preview": preview[:25], "bronze_columns": bronze_cols,
            "silver_columns": silver_cols,
            "reading": self.last_reading, "log": log,
            "headline": f"Silver ← {n} cleaned rows, {n_tags} tags",
            "steps_done": self.steps_done,
        }

    # ---- hop 5: DuckDB + MinIO Gold lake ----------------------------------
    def _step_gold_lake(self) -> Dict:
        if not self.silver_path:
            raise ValueError("Run the Silver transform first (hop 4).")
        batch = self.last_kafka_batch or []
        first_off = batch[0]["offset"] if batch else 0
        last_off = batch[-1]["offset"] if batch else 0
        day = _now_iso()[:10]
        rel = Path("gold") / "telemetry" / f"dt={day}" / f"batch-{first_off}-{last_off}.parquet"
        dest = LAKE_ROOT / rel
        s3_uri = f"s3://mosaic-lake/{rel.as_posix()}"
        silver_uri = _duck_path(Path(self.silver_path))
        dest.parent.mkdir(parents=True, exist_ok=True)
        gold_sql = f"""
CREATE OR REPLACE TABLE gold AS
SELECT
    tag,
    param,
    asset,
    value,
    unit,
    ts,
    quality,
    uns_topic,
    source,
    bronze_offset,
    'reading'           AS grain,
    'contextualize'     AS ready_for,
    'gold'              AS medallion_zone,
    'pending'           AS context_status
FROM read_parquet('{silver_uri}')
"""
        snapshot_sql = """
SELECT tag, param, asset, value, unit, ts, quality, bronze_offset
FROM gold
QUALIFY ROW_NUMBER() OVER (PARTITION BY tag ORDER BY ts DESC NULLS LAST) = 1
"""
        import duckdb
        con = duckdb.connect()
        con.execute(gold_sql)
        con.execute(f"COPY gold TO '{_duck_path(dest)}' (FORMAT PARQUET)")
        preview = self._fetch_dicts(con, "SELECT * FROM gold")
        snapshot = self._fetch_dicts(con, snapshot_sql)
        n = con.execute("SELECT COUNT(*) FROM gold").fetchone()[0]
        n_tags = con.execute("SELECT COUNT(DISTINCT tag) FROM gold").fetchone()[0]
        con.close()
        with self._lock:
            self.gold_path = str(dest)
            for row in preview:
                self.gold_lake.append(row)
            self.steps_done = max(self.steps_done, 5)
        log = [
            "[duckdb]  curate Silver → Gold lake table  (NOT Flink contextualized Gold)",
            f"[duckdb]  READ  silver  {silver_uri}",
            "[contract] grain=reading  ready_for=contextualize  context_status=pending",
            "[contract] columns Flink will consume: tag, param, asset, value, unit, ts, quality, bronze_offset",
            f"[duckdb]  {gold_sql.strip()}",
            f"[duckdb]  latest-per-tag snapshot — {len(snapshot)} tags",
            f"[minio]   PUT  {s3_uri}  ({n} rows)",
            "[role]    Fit-for-purpose store. Layer 3 adds meaning (batch, product, spec, OVER/UNDER).",
            "[role]    Do not confuse this lake Gold with STORE.gold contextualized events.",
            f"✓ gold lake  {n} rows / {n_tags} tags  {s3_uri}  — ready for Contextualize",
        ]
        return {
            "step": 5, "ok": True, "offset": last_off, "zone": "gold-lake",
            "store": "duckdb+minio", "path": str(dest), "s3_uri": s3_uri, "count": n,
            "preview": preview[:25], "snapshot": snapshot,
            "reading": self.last_reading, "log": log,
            "headline": f"Gold lake ← {n} rows, {n_tags} tags ready for Contextualize",
            "steps_done": self.steps_done,
            "note": "Medallion Gold lake — not Layer 3 Flink Gold events",
        }

    def gold_lake_rows(self, n: int = 5000) -> List[Dict]:
        """All Gold-lake readings (parquet if present, else in-memory)."""
        with self._lock:
            path = self.gold_path
            mem = list(self.gold_lake)
        if path and Path(path).exists():
            import duckdb
            con = duckdb.connect()
            rows = self._fetch_dicts(
                con, f"SELECT * FROM read_parquet('{_duck_path(Path(path))}')"
            )
            con.close()
            return rows if not n else rows[-n:]
        return mem[-n:] if n else mem

    def gold_lake_handoff(self) -> Dict:
        """Compact payload Layer 3 uses to start mapping."""
        rows = self.gold_lake_rows()
        latest: Dict[str, Dict] = {}
        for r in rows:
            tag = r.get("tag")
            if tag:
                latest[str(tag)] = r
        st = self.status()
        return {
            **st,
            "ready": bool(rows),
            "count": len(rows),
            "tags": len(latest),
            "snapshot": list(latest.values()),
            "preview": rows[:12],
        }

    def excel_bytes(self, zone: str) -> tuple[bytes, str]:
        """Workbook of the current Medallion parquet for bronze, silver or gold."""
        zone = (zone or "").strip().lower()
        mapping = {
            "bronze": (self.bronze_path, "mosaic-bronze.xlsx", "Bronze"),
            "silver": (self.silver_path, "mosaic-silver.xlsx", "Silver"),
            "gold": (self.gold_path, "mosaic-gold.xlsx", "Gold"),
        }
        if zone not in mapping:
            raise ValueError("zone must be bronze, silver or gold")
        path, filename, sheet = mapping[zone]
        if not path:
            raise ValueError(f"Run the {sheet} hop first so there is a lake table to export")
        import duckdb
        from io import BytesIO
        from openpyxl import Workbook
        con = duckdb.connect()
        rows = self._fetch_dicts(
            con, f"SELECT * FROM read_parquet('{_duck_path(Path(path))}')"
        )
        con.close()
        if not rows:
            raise ValueError(f"{sheet} parquet has no rows")
        wb = Workbook()
        ws = wb.active
        ws.title = sheet[:31]
        cols = list(rows[0].keys())
        ws.append(cols)
        for rec in rows:
            ws.append([_jsonable(rec.get(c)) for c in cols])
        buf = BytesIO()
        wb.save(buf)
        return buf.getvalue(), filename

    @staticmethod
    def _register_row(con, name: str, row: Dict) -> None:
        cols: List[str] = []
        types: List[str] = []
        vals: List[Any] = []
        for k, v in row.items():
            cols.append(k)
            vals.append(v)
            if isinstance(v, bool):
                types.append("BOOLEAN")
            elif isinstance(v, int) and not isinstance(v, bool):
                types.append("BIGINT")
            elif isinstance(v, float):
                types.append("DOUBLE")
            else:
                types.append("VARCHAR")
        schema = ", ".join(f"{c} {t}" for c, t in zip(cols, types))
        placeholders = ", ".join(["?" for _ in cols])
        con.execute(f"CREATE OR REPLACE TABLE {name} ({schema})")
        con.execute(
            f"INSERT INTO {name} ({', '.join(cols)}) VALUES ({placeholders})",
            vals,
        )

    @classmethod
    def _register_rows(cls, con, name: str, rows: List[Dict]) -> None:
        if not rows:
            raise ValueError("no rows to register")
        cls._register_row(con, name, rows[0])
        if len(rows) == 1:
            return
        cols = list(rows[0].keys())
        placeholders = ", ".join(["?" for _ in cols])
        for row in rows[1:]:
            con.execute(
                f"INSERT INTO {name} ({', '.join(cols)}) VALUES ({placeholders})",
                [row.get(c) for c in cols],
            )

    @staticmethod
    def _fetch_dicts(con, sql: str) -> List[Dict]:
        cur = con.execute(sql)
        cols = [c[0] for c in cur.description]
        return [_row_jsonable(dict(zip(cols, row))) for row in cur.fetchall()]


STORE = Store()
